"""知识文档清洗与结构恢复测试。"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document

from core.knowledge.ingestion import (
    BlockType,
    ChunkingConfig,
    DocumentPreprocessor,
    StructuredDocumentChunker,
)
from core.knowledge.ingestion.cleaning import normalize_text, repair_soft_line_breaks


def test_text_cleaning_is_conservative() -> None:
    """清理不可见字符和错误换行，但不破坏型号、单位与版本号。"""

    assert normalize_text("iPhone\u200b 16\tPro  20 W") == "iPhone 16 Pro 20 W"
    assert repair_soft_line_breaks(["支持 USB-C 快速充", "电，最大功率 20 W。"])
    assert repair_soft_line_breaks(["inter-", "national warranty"]) == "international warranty"


def test_markdown_recovers_heading_paths_lists_and_tables(tmp_path: Path) -> None:
    source = tmp_path / "manual.md"
    source.write_text(
        """# iPhone 16 使用说明

<!-- 这段注释不应入库 -->
## 电池与充电

使用兼容的 USB-C 电源适配器
可以快速充电。

- 不要挤压电池
- 发现鼓包时停止使用

表 1：充电规格
| 模式 | 功率 |
| --- | --- |
| 有线 | 20 W |
| 无线 | 15 W |
""",
        encoding="utf-8",
    )

    parsed = DocumentPreprocessor().parse(source, document_id="iphone_16_manual")

    assert parsed.title == "iPhone 16 使用说明"
    assert [block.block_type for block in parsed.blocks] == [
        BlockType.HEADING,
        BlockType.HEADING,
        BlockType.PARAGRAPH,
        BlockType.LIST,
        BlockType.TABLE,
    ]
    paragraph = parsed.blocks[2]
    assert paragraph.section_path == "iPhone 16 使用说明 > 电池与充电"
    assert "适配器可以快速充电" in paragraph.text
    table = parsed.blocks[-1]
    assert table.table is not None
    assert table.table.title == "充电规格"
    assert table.table.headers == ["模式", "功率"]
    assert table.table.row_number_end == 2


def test_docx_preserves_document_order_headings_lists_and_tables(tmp_path: Path) -> None:
    source = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("ThinkPad X1 使用说明", level=1)
    document.add_heading("接口说明", level=2)
    document.add_paragraph("支持 USB-C\u200b 充电。")
    document.add_paragraph("请使用原装适配器。", style="List Bullet")
    document.add_paragraph("连接电源后再开机。", style="List Bullet")
    document.add_paragraph("表 1：接口规格", style="Caption")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "接口"
    table.cell(0, 1).text = "能力"
    table.cell(1, 0).text = "USB-C"
    table.cell(1, 1).text = "充电与数据"
    table.cell(2, 0).text = "HDMI"
    table.cell(2, 1).text = "视频输出"
    document.save(source)

    parsed = DocumentPreprocessor().parse(source)

    assert [block.block_type for block in parsed.blocks] == [
        BlockType.HEADING,
        BlockType.HEADING,
        BlockType.PARAGRAPH,
        BlockType.LIST,
        BlockType.TABLE,
    ]
    assert parsed.blocks[2].text == "支持 USB-C 充电。"
    assert parsed.blocks[2].section_path == "ThinkPad X1 使用说明 > 接口说明"
    assert parsed.blocks[-1].table is not None
    assert parsed.blocks[-1].table.title == "接口规格"
    assert parsed.metadata["pagination_reliable"] is False
    assert parsed.warnings


class _FakeTable:
    def __init__(self) -> None:
        self.bbox = (20.0, 150.0, 400.0, 230.0)

    def extract(self):
        return [["模式", "功率"], ["有线", "20 W"], ["无线", "15 W"]]


class _FakePage:
    height = 800.0

    def __init__(self, page_number: int, words: list[dict], table: bool = False) -> None:
        self.page_number = page_number
        self._words = words
        self._tables = [_FakeTable()] if table else []

    def extract_words(self, **_kwargs):
        return list(self._words)

    def find_tables(self):
        return list(self._tables)


class _FakePdf:
    def __init__(self, pages: list[_FakePage]) -> None:
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False


def _words(text: str, top: float, *, size: float = 10.0, font: str = "Regular") -> list[dict]:
    """构造一行带位置和字体属性的 pdfplumber 单词。"""

    output = []
    x = 30.0
    for word in text.split(" "):
        width = max(18.0, len(word) * 8.0)
        output.append(
            {
                "text": word,
                "x0": x,
                "x1": x + width,
                "top": top,
                "bottom": top + size + 2,
                "size": size,
                "fontname": font,
            }
        )
        x += width + 6
    return output


def test_pdf_removes_repeated_margins_and_recovers_paragraphs_and_tables(
    tmp_path: Path,
) -> None:
    source = tmp_path / "manual.pdf"
    document = pymupdf.open()
    document.new_page(width=600, height=800)
    document.new_page(width=600, height=800)
    first, second = document[0], document[1]
    for page, number in ((first, 1), (second, 2)):
        page.insert_text((50, 25), "Device Manual", fontsize=9)
        page.insert_text((285, 785), str(number), fontsize=8)
    first.insert_text((50, 100), "1. Battery and charging", fontname="hebo", fontsize=14)
    first.insert_text((50, 130), "Use a compatible USB-C adapter for fast charging.")
    second.insert_text((50, 120), "Table 1: Charging specification")

    # 绘制真实表格线，让 pdfplumber 负责有边框表格提取。
    x_positions = (50, 200, 350)
    y_positions = (150, 180, 210, 240)
    for x in x_positions:
        second.draw_line((x, y_positions[0]), (x, y_positions[-1]))
    for y in y_positions:
        second.draw_line((x_positions[0], y), (x_positions[-1], y))
    cells = (
        ("Mode", "Power"),
        ("Wired", "20 W"),
        ("Wireless", "15 W"),
    )
    for row, values in enumerate(cells):
        for column, value in enumerate(values):
            second.insert_text((x_positions[column] + 5, y_positions[row] + 20), value)
    document.save(source)
    document.close()

    parsed = DocumentPreprocessor().parse(source)

    all_text = "\n".join(block.text for block in parsed.blocks)
    assert "\n1\n" not in f"\n{all_text}\n"
    assert all_text.count("Device Manual") == 0
    assert "compatible USB-C adapter" in all_text
    table = next(block for block in parsed.blocks if block.block_type == BlockType.TABLE)
    assert table.page_start == table.page_end == 2
    assert table.table is not None
    assert table.table.title == "Charging specification"
    assert table.table.row_number_end == 2


def test_chunker_respects_sections_and_keeps_semantic_overlap(tmp_path: Path) -> None:
    source = tmp_path / "long_manual.md"
    source.write_text(
        """# 手机说明书
## 充电
第一段介绍 USB-C 充电协议和适配器要求。

第二段说明最大功率以及充电时的温度限制。

第三段提醒用户不要使用损坏的数据线。

## 安全
电池鼓包时应立即停止使用并远离热源。
""",
        encoding="utf-8",
    )
    parsed = DocumentPreprocessor().parse(source, document_id="phone_manual")
    chunker = StructuredDocumentChunker(
        ChunkingConfig(target_tokens=30, max_tokens=45, overlap_tokens=12)
    )

    chunks = chunker.chunk(
        parsed,
        document_type="product_manual",
        product_model="demo_phone",
        category="smartphone",
        version="v1",
    )

    assert len(chunks) >= 3
    assert chunks[0].section_path == "手机说明书 > 充电"
    assert any(chunk.section == "安全" for chunk in chunks)
    assert all(chunk.token_count <= 60 for chunk in chunks)
    assert all(chunk.chunk_id.startswith("phone_manual__") for chunk in chunks)
    assert len({chunk.chunk_id for chunk in chunks}) == len(chunks)
    assert chunks[0].to_langchain_document().metadata["product_model"] == "demo_phone"


def test_chunker_splits_large_tables_by_complete_rows(tmp_path: Path) -> None:
    source = tmp_path / "table.md"
    rows = "\n".join(f"| 型号{i} | USB-C | {i + 10} W |" for i in range(12))
    source.write_text(
        "# 充电规格\n\n## 型号对照\n\n表 2：型号功率表\n"
        "| 型号 | 接口 | 最大功率 |\n| --- | --- | --- |\n"
        + rows,
        encoding="utf-8",
    )
    parsed = DocumentPreprocessor().parse(source, document_id="table_manual")
    chunks = StructuredDocumentChunker(
        ChunkingConfig(target_tokens=50, max_tokens=65, overlap_tokens=10)
    ).chunk(parsed, document_type="product_manual", category="smartphone")

    table_chunks = [chunk for chunk in chunks if chunk.chunk_type == "table"]
    assert len(table_chunks) > 1
    assert all("型号功率表" in chunk.content for chunk in table_chunks)
    assert all("| 行号 | 型号 | 接口 | 最大功率 |" in chunk.content for chunk in table_chunks)
    assert table_chunks[0].table_row_start == 1
    assert table_chunks[-1].table_row_end == 12
    assert all(chunk.table_row_start <= chunk.table_row_end for chunk in table_chunks)


def test_chunker_merges_short_intro_into_following_table(tmp_path: Path) -> None:
    source = tmp_path / "intro_table.md"
    source.write_text(
        """# 路由器说明书

## 产品概览

以下参数用于识别本型号，相似型号不得混用。

| 项目 | 参数 |
| --- | --- |
| 硬件型号 | RT-X2 |
| 供电 | 12 V 2 A |
""",
        encoding="utf-8",
    )
    parsed = DocumentPreprocessor().parse(source, document_id="router_manual")
    chunks = StructuredDocumentChunker().chunk(
        parsed,
        document_type="product_manual",
        product_model="rt_x2",
        category="router",
    )

    table_chunks = [chunk for chunk in chunks if chunk.chunk_type == "table"]
    assert len(table_chunks) == 1
    assert "以下参数用于识别本型号，相似型号不得混用。" in table_chunks[0].content
    assert "| 硬件型号 | RT-X2 |" in table_chunks[0].content
    assert not any(
        chunk.chunk_type == "text" and "以下参数用于识别本型号" in chunk.content
        for chunk in chunks
    )


def test_large_table_chunks_overlap_by_two_complete_rows(tmp_path: Path) -> None:
    source = tmp_path / "overlapping_table.md"
    rows = "\n".join(f"| 型号{i} | USB-C | {i + 10} W |" for i in range(20))
    source.write_text(
        "# 充电规格\n\n## 型号对照\n\n以下是各型号的充电参数。\n\n"
        "| 型号 | 接口 | 最大功率 |\n| --- | --- | --- |\n"
        + rows,
        encoding="utf-8",
    )
    parsed = DocumentPreprocessor().parse(source, document_id="overlap_table_manual")
    chunks = StructuredDocumentChunker(
        ChunkingConfig(target_tokens=100, max_tokens=150, overlap_tokens=10)
    ).chunk(parsed, document_type="product_manual", category="smartphone")

    table_chunks = [chunk for chunk in chunks if chunk.chunk_type == "table"]
    assert len(table_chunks) > 1
    assert all("以下是各型号的充电参数。" in chunk.content for chunk in table_chunks)
    assert all(
        current.table_row_start == previous.table_row_end - 1
        for previous, current in zip(table_chunks, table_chunks[1:])
    )
    assert all(chunk.token_count <= 150 for chunk in table_chunks)


def test_pdf_recovers_borderless_table_from_bold_header_and_aligned_columns(
    tmp_path: Path,
) -> None:
    source = tmp_path / "borderless.pdf"
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.insert_text((50, 70), "Robot X2 Manual", fontname="hebo")
    page.insert_text((50, 100), "1. Indicator status", fontname="hebo")
    for x, value in zip((50, 180, 320), ("Status", "Meaning", "Action"), strict=True):
        page.insert_text((x, 140), value, fontname="hebo")
    for y, values in (
        (170, ("Red5", "LidarBlocked", "CleanLidar")),
        (195, ("Yellow", "LowBattery", "Recharge")),
    ):
        for x, value in zip((50, 180, 320), values, strict=True):
            page.insert_text((x, y), value)
    page.insert_text((50, 235), "2. Warranty", fontname="hebo")
    document.save(source)
    document.close()

    parsed = DocumentPreprocessor().parse(source)

    table = next(block for block in parsed.blocks if block.block_type == BlockType.TABLE)
    assert table.table is not None
    assert table.table.headers == ["Status", "Meaning", "Action"]
    assert table.table.rows[0] == ["Red5", "LidarBlocked", "CleanLidar"]
    assert table.section == "1. Indicator status"


def test_chunker_does_not_copy_an_oversized_block_as_overlap(tmp_path: Path) -> None:
    source = tmp_path / "overlap.md"
    source.write_text(
        "# 说明书\n\n## 测试\n\n" + "甲" * 20 + "\n\n" + "乙" * 20,
        encoding="utf-8",
    )
    parsed = DocumentPreprocessor().parse(source, document_id="overlap_manual")
    chunks = StructuredDocumentChunker(
        ChunkingConfig(target_tokens=30, max_tokens=50, overlap_tokens=10)
    ).chunk(parsed, document_type="product_manual")

    assert len(chunks) == 2
    assert "甲" in chunks[0].content
    assert "甲" not in chunks[1].content
    assert "乙" in chunks[1].content
