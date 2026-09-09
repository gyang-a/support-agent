"""生成用于文档清洗与证据级切片评测的多格式合成知识库。"""

from __future__ import annotations

import hashlib
import json
import shutil
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from pypdf import PdfReader, PdfWriter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = ROOT / "data" / "knowledge" / "retrieval_eval_v1"
DOCUMENT_ROOT = OUTPUT_ROOT / "documents"
VERSION = "2026.08-v2"


PRODUCTS: list[dict[str, Any]] = [
    {
        "id": "aurora_x1", "model": "Aurora X1", "format": "docx", "category": "smartphone",
        "sibling": "aurora_x1_pro", "identity": "6.10 英寸 OLED，机身识别码前缀 AX1-CN",
        "power": "最大有线充电功率为 45 W，仅 USB-C1 口支持快充",
        "temperature": "允许充电环境温度为 0 至 35 摄氏度",
        "setup": "首次开机后长按电源键 3 秒，再扫描包装内的设备迁移二维码",
        "feature": "护眼阅读模式只能在刷新率设置为 60 Hz 时启用",
        "indicator": "出现 E17 且充电图标闪烁时，表示 USB-C 端口检测到液体",
        "reset": "同时按住电源键和音量减键 12 秒可执行强制重启",
        "compat": "支持 Wi-Fi 6，不支持 Wi-Fi 7 的 320 MHz 信道",
        "warranty": "电池在 12 个月内且健康度低于 80% 时可申请检测，不保证直接换新",
        "safety": "端口进液后至少静置 4 小时，禁止使用热风机烘干",
    },
    {
        "id": "aurora_x1_pro", "model": "Aurora X1 Pro", "format": "pdf", "category": "smartphone",
        "sibling": "aurora_x1", "identity": "6.72 英寸 LTPO OLED，机身识别码前缀 AXP-CN",
        "power": "最大有线充电功率为 65 W，USB-C1 和 USB-C2 均支持快充",
        "temperature": "允许充电环境温度为 5 至 38 摄氏度",
        "setup": "首次开机后长按电源键 2 秒，并使用星澜账号完成安全校验",
        "feature": "专业影像模式要求剩余电量不低于 20%，并关闭超级省电",
        "indicator": "出现 E19 且状态灯呈蓝红交替时，表示镜头防抖模块需要校准",
        "reset": "同时按住电源键和音量加键 15 秒可执行强制重启",
        "compat": "支持 Wi-Fi 7，但 320 MHz 信道仅在 6 GHz 频段可用",
        "warranty": "潜望镜模组享受 24 个月有限保修，人为跌落造成的偏轴不在范围内",
        "safety": "进行镜头校准时必须平放设备，禁止在移动交通工具内执行",
    },
    {
        "id": "novabook_14", "model": "NovaBook 14", "format": "docx", "category": "laptop",
        "sibling": "novabook_14_pro", "identity": "14.0 英寸 1920x1200 屏幕，产品代码 NB14-26",
        "power": "随机电源适配器为 65 W，只有左侧后方 USB-C 口支持整机充电",
        "temperature": "建议在 10 至 32 摄氏度环境中为电池充电",
        "setup": "首次启动连接电源后，按电源键一次并等待键盘背光连续亮起",
        "feature": "静音模式将处理器持续功耗限制在 18 W，适合会议场景",
        "indicator": "电源灯连续三次橙色短闪表示适配器功率不足 45 W",
        "reset": "关机后按住电源键 10 秒，松开 5 秒后再次短按可重置电源控制器",
        "compat": "支持 DisplayPort 1.4 输出，不支持通过右侧 USB-C 口连接显示器",
        "warranty": "固态硬盘有限保修为 24 个月，自行更换内存不会自动取消整机保修",
        "safety": "清洁键盘前必须关机并断开电源，禁止将清洁剂直接喷向键帽",
    },
    {
        "id": "novabook_14_pro", "model": "NovaBook 14 Pro", "format": "pdf", "category": "laptop",
        "sibling": "novabook_14", "identity": "14.5 英寸 2880x1800 OLED 屏幕，产品代码 NBP14-26",
        "power": "随机电源适配器为 100 W，左右两个 USB-C4 端口均可为整机充电",
        "temperature": "建议在 5 至 35 摄氏度环境中为电池充电",
        "setup": "首次启动必须连接 100 W 适配器，并保持电源键按下 2 秒",
        "feature": "创作模式允许处理器与显卡合计功耗达到 72 W，但必须接通电源",
        "indicator": "电源灯连续五次白色短闪表示 BIOS 恢复映像校验失败",
        "reset": "关机后同时按住 Fn 和电源键 15 秒可重置嵌入式控制器",
        "compat": "两个 USB-C4 端口都支持 DisplayPort 2.1，HDMI 口最高输出 4K 120 Hz",
        "warranty": "OLED 面板亮点政策适用 12 个月，静态画面造成的烧屏不属于亮点故障",
        "safety": "高性能运行时不得遮挡转轴出风口，并应保留至少 10 厘米散热空间",
    },
    {
        "id": "sonicpods_air", "model": "SonicPods Air", "format": "docx", "category": "earbuds",
        "sibling": "sonicpods_pro", "identity": "半入耳式耳机，充电盒型号 SPA-C20",
        "power": "充电盒输入上限为 5 V 1 A，不支持无线充电",
        "temperature": "充电盒允许充电温度为 0 至 40 摄氏度",
        "setup": "打开盒盖后按住背面配对键 3 秒，白灯慢闪时选择 SonicPods Air",
        "feature": "游戏低延迟模式通过左右耳同时双击开启，通话中不可启用",
        "indicator": "盒灯红色快速闪烁四次表示左右耳机固件版本不一致",
        "reset": "两只耳机放入盒内并开盖，按住配对键 8 秒直到白灯熄灭再亮起",
        "compat": "支持 AAC 和 SBC，不支持 LDAC 或 LC3 编解码",
        "warranty": "耳机主体有限保修 12 个月，随附耳帽属于耗材不单独保修",
        "safety": "运动后应擦干充电触点再入盒，汗液未干时禁止充电",
    },
    {
        "id": "sonicpods_pro", "model": "SonicPods Pro", "format": "md", "category": "earbuds",
        "sibling": "sonicpods_air", "identity": "入耳式降噪耳机，充电盒型号 SPP-C30",
        "power": "充电盒输入上限为 9 V 2 A，同时支持 Qi 2 W 无线充电",
        "temperature": "充电盒允许充电温度为 5 至 35 摄氏度",
        "setup": "打开盒盖后按住正面触控区 5 秒，紫灯呼吸时选择 SonicPods Pro",
        "feature": "自适应降噪需要完成耳道扫描，并且两只耳机电量都高于 15%",
        "indicator": "盒灯紫色连续闪烁六次表示耳道扫描数据已损坏",
        "reset": "两只耳机放入盒内并合盖 10 秒，开盖后按住触控区 12 秒",
        "compat": "支持 AAC、LDAC 和 LC3，LDAC 与空间音频不能同时开启",
        "warranty": "降噪麦克风有限保修 18 个月，滤网堵塞不属于硬件故障",
        "safety": "耳道有炎症或疼痛时应停止使用，不得使用酒精浸泡耳机主体",
    },
    {
        "id": "meshwave_ax3000", "model": "MeshWave AX3000", "format": "pdf", "category": "router",
        "sibling": "meshwave_ax6000", "identity": "双频 Wi-Fi 6 路由器，硬件版本 MW3-R2",
        "power": "必须使用 12 V 2 A 直流适配器，接口为内正外负",
        "temperature": "工作环境温度为 0 至 40 摄氏度",
        "setup": "主节点接通光猫后等待蓝灯常亮，再使用 MeshWave 应用扫描底部二维码",
        "feature": "最多组建 4 个节点，有线回程时必须连接 LAN2 端口",
        "indicator": "状态灯红色每两秒闪一次表示主节点无法获取上级网络地址",
        "reset": "通电状态下按住 RESET 孔 8 秒，黄灯亮起后松开",
        "compat": "2.4 GHz 支持 WPA2/WPA3 混合模式，访客网络仅支持 WPA2",
        "warranty": "主机有限保修 24 个月，雷击或错误电压造成的损坏不在范围内",
        "safety": "设备四周应保留 8 厘米空间，不得放入弱电箱密闭运行",
    },
    {
        "id": "meshwave_ax6000", "model": "MeshWave AX6000", "format": "docx", "category": "router",
        "sibling": "meshwave_ax3000", "identity": "三频 Wi-Fi 6E 路由器，硬件版本 MW6-R3",
        "power": "必须使用 12 V 3 A 直流适配器，接口为内正外负",
        "temperature": "工作环境温度为 5 至 45 摄氏度",
        "setup": "主节点接通光猫后等待青灯常亮，再使用 MeshWave 应用输入底部八位配对码",
        "feature": "最多组建 6 个节点，6 GHz 专用回程开启后不再广播该频段访客网络",
        "indicator": "状态灯琥珀色每秒闪两次表示 6 GHz 回程信号弱于 -72 dBm",
        "reset": "通电状态下按住 RESET 孔 10 秒，白灯快速闪烁后松开",
        "compat": "支持 WPA3 Enterprise，旧版 WPA 设备只能连接隔离的 IoT 网络",
        "warranty": "主机有限保修 36 个月，第三方固件造成的启动失败不在范围内",
        "safety": "不得叠放两个主节点，设备顶部至少保留 12 厘米散热空间",
    },
    {
        "id": "chargehub_65", "model": "ChargeHub 65", "format": "md", "category": "charger",
        "sibling": "chargehub_100", "identity": "双口氮化镓充电器，型号 CH65-G2",
        "power": "单独使用 USB-C1 时最高 65 W，双口同时使用时分配为 45 W 加 20 W",
        "temperature": "满载工作允许外壳表面温度最高 62 摄氏度",
        "setup": "先连接设备端，再将充电器插入墙上插座，可减少协议重新协商",
        "feature": "低电流模式需长按侧键 2 秒，指示灯变为绿色常亮",
        "indicator": "指示灯红色闪烁三次后熄灭表示输出过流保护已触发",
        "reset": "从插座拔下并断开全部线缆 30 秒即可复位保护状态",
        "compat": "USB-C1 支持 PD 3.0 与 PPS，USB-A 口不支持 PPS",
        "warranty": "充电器主体有限保修 18 个月，第三方线缆不在保修范围内",
        "safety": "不得接入旅行转换器的剃须刀专用口，也不得覆盖充电器散热",
    },
    {
        "id": "chargehub_100", "model": "ChargeHub 100", "format": "docx", "category": "charger",
        "sibling": "chargehub_65", "identity": "三口氮化镓充电器，型号 CH100-G3",
        "power": "单独使用 USB-C1 时最高 100 W，三口同时使用时分配为 65 W、20 W 和 15 W",
        "temperature": "满载工作允许外壳表面温度最高 68 摄氏度",
        "setup": "使用 100 W 输出前应确认线缆带有 5 A E-Marker 芯片",
        "feature": "桌面供电模式需双击侧键，启用后 USB-C1 保留至少 65 W",
        "indicator": "指示灯红白交替闪烁五次表示交流输入电压异常",
        "reset": "从插座拔下并按住侧键 5 秒，再静置 60 秒可复位保护状态",
        "compat": "两个 USB-C 口均支持 PD 3.1，USB-A 口仅支持最高 18 W QC",
        "warranty": "充电器主体有限保修 24 个月，插脚弯曲属于外力损坏",
        "safety": "100 W 满载时必须使用独立墙插，禁止与大功率加热设备共用排插",
    },
]


EXPANSION_FAMILIES = [
    ("luma_s2", "Luma S2", "Luma S2 Pro", "smartphone"),
    ("orbit_fold", "Orbit Fold", "Orbit Fold Max", "smartphone"),
    ("terrabook_13", "TerraBook 13", "TerraBook 13 Plus", "laptop"),
    ("forgebook_16", "ForgeBook 16", "ForgeBook 16 Studio", "laptop"),
    ("echobuds", "EchoBuds Lite", "EchoBuds Max", "earbuds"),
    ("wavebuds", "WaveBuds Sport", "WaveBuds Sport Pro", "earbuds"),
    ("netnest", "NetNest AX1800", "NetNest AX5400", "router"),
    ("skymesh", "SkyMesh BE3600", "SkyMesh BE7200", "router"),
    ("voltdock", "VoltDock 45", "VoltDock 90", "charger"),
    ("powerbrick", "PowerBrick 120", "PowerBrick 160", "charger"),
    ("pulsewatch", "PulseWatch 4", "PulseWatch 4 Ultra", "wearable"),
    ("trailwatch", "TrailWatch S", "TrailWatch S Pro", "wearable"),
    ("canvastab", "CanvasTab 11", "CanvasTab 13", "tablet"),
    ("noteslate", "NoteSlate Mini", "NoteSlate Pro", "tablet"),
    ("lenscraft", "LensCraft C1", "LensCraft C1 Pro", "camera"),
    ("pocketcam", "PocketCam V2", "PocketCam V2 Max", "camera"),
    ("beambox", "BeamBox P1", "BeamBox P1 Plus", "projector"),
    ("cinemapod", "CinemaPod Mini", "CinemaPod 4K", "projector"),
    ("hometone", "HomeTone S", "HomeTone S Max", "speaker"),
    ("airsense", "AirSense Mini", "AirSense Pro", "air_monitor"),
]


def _expansion_facts(
    base_id: str, model: str, category: str, pair_number: int, pro: bool
) -> dict[str, str]:
    """为扩展产品生成成对、可区分且完全虚构的事实。"""

    suffix = "pro" if pro else "standard"
    code = f"{base_id[:3].upper()}{pair_number:02d}-{'P' if pro else 'S'}"
    fast = pair_number + (18 if pro else 8)
    reset_seconds = 8 + pair_number % 5 + (4 if pro else 0)
    warranty_months = 12 + pair_number % 4 * 6 + (12 if pro else 0)
    common = {
        "smartphone": {
            "identity": f"{6.0 + pair_number / 20 + (0.3 if pro else 0):.2f} 英寸 OLED，硬件代码 {code}",
            "power": f"USB-C 主端口最高支持 {35 + fast} W 有线充电，副端口{'支持' if pro else '不支持'}视频输出",
            "temperature": f"建议充电环境温度为 {5 if pro else 0} 至 {38 if pro else 35} 摄氏度",
            "setup": f"长按电源键 {2 if pro else 3} 秒，输入包装标签上的六位迁移码 {pair_number:02d}{pair_number + 11:02d}{pair_number + 23:02d}",
            "feature": f"影像增强模式要求电量高于 {15 + pair_number % 4 * 5}%，并关闭低电量模式",
            "indicator": f"错误码 M{pair_number:02d}{9 if pro else 4} 表示{'相机稳定器' if pro else 'USB-C 端口'}需要重新校准",
            "reset": f"同时按住电源键和音量{'加' if pro else '减'}键 {reset_seconds} 秒执行强制重启",
            "compat": f"支持 Wi-Fi {7 if pro else 6}，{'支持' if pro else '不支持'} 6 GHz 频段",
            "safety": "设备进液后禁止充电或使用热风烘干，应断电并自然通风",
        },
        "laptop": {
            "identity": f"{13 + pair_number % 4 + (0.5 if pro else 0):.1f} 英寸显示屏，主板代码 {code}",
            "power": f"随机适配器为 {65 + (35 if pro else 0) + pair_number % 3 * 10} W，仅标有闪电图标的 USB-C 口支持整机充电",
            "temperature": f"电池充电环境温度为 {8 if pro else 10} 至 {35 if pro else 32} 摄氏度",
            "setup": f"连接原装适配器后短按电源键，并等待键盘背光闪烁 {2 + pair_number % 4} 次",
            "feature": f"性能模式将整机功耗限制设为 {45 + fast} W，必须接通电源才能启用",
            "indicator": f"电源灯连续 {3 + pair_number % 4} 次琥珀色短闪表示适配器握手失败",
            "reset": f"关机后按住 Fn 和电源键 {reset_seconds} 秒可重置嵌入式控制器",
            "compat": f"支持 DisplayPort {2.1 if pro else 1.4}，右侧 USB-C {'支持' if pro else '不支持'}外接显示器",
            "safety": "高负载运行时不得遮挡转轴出风口，清洁前必须关机并断电",
        },
        "earbuds": {
            "identity": f"{'入耳式主动降噪' if pro else '半入耳式'}耳机，充电盒代码 {code}",
            "power": f"充电盒输入上限为 {9 if pro else 5} V {2 if pro else 1} A，{'支持 2 W Qi 无线充电' if pro else '不支持无线充电'}",
            "temperature": f"充电盒允许充电温度为 {5 if pro else 0} 至 {35 if pro else 40} 摄氏度",
            "setup": f"开盖后按住配对区域 {3 + (2 if pro else 0)} 秒，状态灯呈{'紫色呼吸' if pro else '白色慢闪'}时选择设备",
            "feature": f"低延迟模式要求两只耳机电量均高于 {10 + pair_number % 4 * 5}%",
            "indicator": f"盒灯{'紫色' if pro else '红色'}闪烁 {4 + pair_number % 4} 次表示左右耳固件版本不一致",
            "reset": f"耳机入盒并开盖，按住配对区域 {reset_seconds} 秒直到状态灯重新亮起",
            "compat": f"支持 AAC、SBC{' 和 LC3' if pro else ''}，{'支持' if pro else '不支持'}多点连接",
            "safety": "汗液未干时禁止放入充电盒，耳道不适时应立即停止使用",
        },
        "router": {
            "identity": f"{'三频' if pro else '双频'}无线网状路由器，硬件版本 {code}",
            "power": f"必须使用 12 V {3 if pro else 2} A 直流适配器，接口为内正外负",
            "temperature": f"工作环境温度为 {5 if pro else 0} 至 {45 if pro else 40} 摄氏度",
            "setup": f"连接上级网络后等待{'青' if pro else '蓝'}灯常亮，再输入机身底部配对码 {pair_number + 3200}",
            "feature": f"最多组建 {6 if pro else 4} 个节点，有线回程必须连接 LAN{2 + pair_number % 2} 端口",
            "indicator": f"状态灯{'琥珀色' if pro else '红色'}每秒闪烁 {2 if pro else 1} 次表示回程链路异常",
            "reset": f"通电后按住 RESET 孔 {reset_seconds} 秒，状态灯快速闪烁后松开",
            "compat": f"支持 WPA3{' Enterprise' if pro else ''}，访客网络{'支持 WPA3' if pro else '仅支持 WPA2'}",
            "safety": "设备不得放入密闭弱电箱，顶部与四周必须保留散热空间",
        },
        "charger": {
            "identity": f"{'三口' if pro else '双口'}氮化镓充电器，型号代码 {code}",
            "power": f"USB-C1 单口最高 {70 + fast} W，多口同时使用时主口至少保留 {40 + fast} W",
            "temperature": f"满载工作允许外壳表面温度最高 {58 + pair_number % 5 + (5 if pro else 0)} 摄氏度",
            "setup": f"使用高功率输出前确认线缆具有 {5 if pro else 3} A E-Marker 标识",
            "feature": f"低电流模式需长按侧键 {2 + pair_number % 3} 秒，绿灯常亮后生效",
            "indicator": f"指示灯红白交替闪烁 {3 + pair_number % 4} 次表示输入电压异常",
            "reset": f"拔下全部线缆并静置 {30 + pair_number * 2} 秒可复位保护状态",
            "compat": f"USB-C1 支持 PD {3.1 if pro else 3.0} 与 PPS，USB-A 最高 {18 + pair_number % 3 * 4} W",
            "safety": "满载时必须使用独立墙插，不得覆盖散热面或与加热设备共用排插",
        },
        "wearable": {
            "identity": f"{42 + pair_number % 4 + (3 if pro else 0)} 毫米智能手表，设备代码 {code}",
            "power": f"磁吸底座输入为 5 V {1.5 if pro else 1} A，{'支持反向无线补电' if pro else '不支持通用 Qi 充电'}",
            "temperature": f"允许充电温度为 {5 if pro else 0} 至 {38 if pro else 35} 摄氏度",
            "setup": f"长按侧键 {2 + pair_number % 3} 秒，扫描表盘显示的绑定二维码",
            "feature": f"连续血氧模式要求佩戴紧度为 {2 + pair_number % 2} 级且电量高于 20%",
            "indicator": f"屏幕显示 H{pair_number:02d} 表示光学传感器需要清洁并重新贴合",
            "reset": f"同时按住侧键和功能键 {reset_seconds} 秒，徽标出现后松开",
            "compat": f"支持 Android {12 + pair_number % 3} 及以上，{'支持' if pro else '不支持'}独立 eSIM",
            "safety": "皮肤出现红肿时应停止佩戴，充电前必须擦干表背和底座",
        },
        "tablet": {
            "identity": f"{10 + pair_number % 3 + (1.5 if pro else 0):.1f} 英寸平板，硬件代码 {code}",
            "power": f"USB-C 端口最高支持 {30 + fast} W PD 充电，{'支持' if pro else '不支持'}外接显示器供电",
            "temperature": f"建议充电环境温度为 {5 if pro else 0} 至 {36 if pro else 34} 摄氏度",
            "setup": f"开机后连接 Wi-Fi，并输入随附卡片上的初始化码 T{pair_number:02d}{pair_number + 44:02d}",
            "feature": f"手写低延迟模式要求刷新率设为 {120 if pro else 90} Hz",
            "indicator": f"状态栏出现 P{pair_number:02d} 表示触控笔固件校验未通过",
            "reset": f"按住电源键和音量减键 {reset_seconds} 秒可执行强制重启",
            "compat": f"支持 USB {3.2 if pro else 2.0}，{'支持' if pro else '不支持'} DisplayPort Alt Mode",
            "safety": "屏幕破裂时禁止继续使用触控笔，设备弯曲后不得自行压平",
        },
        "camera": {
            "identity": f"{'全画幅' if pro else 'APS-C'}无反相机，机身代码 {code}",
            "power": f"USB-C 供电最高 {30 + fast} W，连续录像需使用标有 PD 的端口",
            "temperature": f"连续录像环境温度建议为 {5 if pro else 0} 至 {40 if pro else 35} 摄氏度",
            "setup": f"装入电池后长按菜单键 {2 + pair_number % 3} 秒，再执行日期和传感器初始化",
            "feature": f"高码率录像要求存储卡持续写入不低于 {120 + pair_number * 8 + (80 if pro else 0)} MB/s",
            "indicator": f"机身灯连续闪烁 C{pair_number:02d} 节奏表示传感器防抖需要校准",
            "reset": f"关机取下电池 {20 + pair_number} 秒，再按住快门键装回电池",
            "compat": f"支持 UHS-{'II' if pro else 'I'}，{'支持' if pro else '不支持'}外录 RAW 视频",
            "safety": "清洁传感器时镜头卡口必须朝下，禁止使用罐装高压气体",
        },
        "projector": {
            "identity": f"{'4K' if pro else '1080p'} 便携投影仪，光机代码 {code}",
            "power": f"必须使用 {90 + fast} W 适配器，USB-C 供电不足时自动限制亮度",
            "temperature": f"工作环境温度为 {5 if pro else 0} 至 {35 if pro else 32} 摄氏度",
            "setup": f"放置在水平桌面后长按电源键 {2 + pair_number % 3} 秒并执行梯形校正",
            "feature": f"影院模式要求投射距离不少于 {1.5 + pair_number / 20:.1f} 米并关闭节能模式",
            "indicator": f"温度灯闪烁 {3 + pair_number % 4} 次表示进风口阻塞",
            "reset": f"断电静置 {60 + pair_number * 3} 秒，再按住电源键 {reset_seconds} 秒恢复",
            "compat": f"HDMI 支持最高 {'4K 60 Hz' if pro else '1080p 120 Hz'}，USB 口不支持视频输入",
            "safety": "运行时不得直视镜头或遮挡进出风口，关机后应等待风扇停止",
        },
        "speaker": {
            "identity": f"{'双单元' if pro else '单单元'}智能音箱，设备代码 {code}",
            "power": f"电源输入为 {20 + fast} W USB-C，低于额定功率时禁用高音量模式",
            "temperature": f"工作环境温度为 {0 if pro else 5} 至 {40 if pro else 35} 摄氏度",
            "setup": f"接通电源后按住麦克风键 {3 + pair_number % 3} 秒并在应用中输入配对码 {pair_number + 6100}",
            "feature": f"立体声组网要求两台设备固件版本一致且距离不超过 {4 + pair_number % 3} 米",
            "indicator": f"环形灯橙色旋转 {3 + pair_number % 4} 圈表示网络认证失败",
            "reset": f"同时按住音量加减键 {reset_seconds} 秒直到白灯常亮",
            "compat": f"支持 Wi-Fi {6 if pro else 5} 和蓝牙 {5.3 if pro else 5.0}，不支持作为 USB 声卡",
            "safety": "不得放置在水槽旁或用湿布覆盖麦克风孔，雷雨时应断开电源",
        },
        "air_monitor": {
            "identity": f"桌面空气质量监测仪，传感器模块代码 {code}",
            "power": f"USB-C 输入为 5 V {2 if pro else 1} A，断电后数据缓存可维持 {20 + pair_number} 分钟",
            "temperature": f"传感器校准环境温度为 {10 if pro else 15} 至 {30 if pro else 28} 摄氏度",
            "setup": f"通电预热 {5 + pair_number % 4} 分钟，再按住校准键 3 秒完成基线设置",
            "feature": f"自动校准要求连续运行 {5 + pair_number % 3} 天且每日通风不少于 30 分钟",
            "indicator": f"屏幕显示 S{pair_number:02d} 表示颗粒物传感器风道堵塞",
            "reset": f"断开电源并按住校准键 {reset_seconds} 秒，重新通电后释放",
            "compat": f"支持 Wi-Fi {6 if pro else 4}，{'支持' if pro else '不支持'}本地 MQTT 数据导出",
            "safety": "不得向进气口喷洒清洁剂或香水，清理滤网前必须断电",
        },
    }[category]
    return {
        **common,
        "warranty": f"主体有限保修 {warranty_months} 个月，非授权拆机和错误供电不在范围内",
        "id_suffix": suffix,
    }


def _build_expansion_products() -> list[dict[str, Any]]:
    formats = ("docx", "pdf", "md", "docx", "pdf", "docx", "md")
    products: list[dict[str, Any]] = []
    for pair_number, (base_id, standard_model, pro_model, category) in enumerate(
        EXPANSION_FAMILIES, start=1
    ):
        standard_id = f"{base_id}_standard"
        pro_id = f"{base_id}_pro"
        for pro, document_id, model, sibling in (
            (False, standard_id, standard_model, pro_id),
            (True, pro_id, pro_model, standard_id),
        ):
            product = {
                "id": document_id,
                "model": model,
                "format": formats[len(products) % len(formats)],
                "category": category,
                "sibling": sibling,
            }
            product.update(_expansion_facts(base_id, model, category, pair_number, pro))
            products.append(product)
    return products


def _enrich_product(product: dict[str, Any], index: int) -> None:
    noun = {
        "smartphone": "移动设备", "laptop": "笔记本电脑", "earbuds": "耳机",
        "router": "网络设备", "charger": "供电设备", "wearable": "穿戴设备",
        "tablet": "平板设备", "camera": "影像设备", "projector": "显示设备",
        "speaker": "音频设备", "air_monitor": "监测设备",
    }[product["category"]]
    product["spec_detail"] = (
        f"{noun}运行日志最多保留 {14 + index % 17} 天，导出文件前缀为 "
        f"{product['id'].upper()}-LOG"
    )
    product["daily_use"] = (
        f"日常使用时每 {20 + index % 8 * 5} 分钟自动保存一次配置，切换模式前应等待保存图标消失"
    )
    product["maintenance"] = (
        f"建议每 {30 + index % 6 * 15} 天检查接口、通风位置和附件连接状态，并清除表面灰尘"
    )
    product["backup"] = (
        f"恢复出厂设置前应导出配置包，文件校验码前缀为 BK{index + 101:03d}"
    )
    product["indicator_secondary"] = (
        f"状态灯黄色长亮 {2 + index % 5} 秒表示本地配置写入尚未完成，此时不得断电"
    )
    product["diagnostic"] = (
        f"若连续 {2 + index % 4} 次启动失败，应断开外设并记录诊断编号 D{index + 210:03d} 后联系支持"
    )
    product["service"] = (
        f"送修时需同时提供设备序列号、购买凭证和最近一次日志包，服务受理代码为 SV{index + 500:03d}"
    )


PRODUCTS.extend(_build_expansion_products())
for _product_index, _product in enumerate(PRODUCTS):
    _enrich_product(_product, _product_index)


SECTIONS = [
    "1 产品概览",
    "2 初次设置与功能",
    "3 供电与兼容性",
    "4 维护与数据",
    "5 故障诊断",
    "6 保修与服务",
    "7 安全要求",
]


def _font(run, size: float = 11, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_margins(cell: Any, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _style_table(table: Any, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            if row_index == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E8EEF5")
                cell._tc.get_or_add_tcPr().append(shd)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _font(run, 9.5, bold=row_index == 0)


def _add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run("第 ")
    _font(run, 9, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    _font(run, 9, color="666666")


def _configure_docx(doc: Document, product: dict[str, Any]) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = f"星澜数字设备支持中心 | {product['model']} 用户指南"
    for run in header.runs:
        _font(run, 9, color="667085")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(f"内部评测样本 | {VERSION} | ")
    _font(run, 9, color="667085")
    _add_page_field(footer)


def _add_docx_title(doc: Document, product: dict[str, Any]) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    _font(kicker.add_run("产品使用与服务指南"), 10, True, "7A5A00")
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _font(title.add_run(product["model"]), 28, True, "0B2545")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    _font(subtitle.add_run(f"版本 {VERSION} | 合成评测文档，不代表真实商品规格"), 10, False, "667085")


def _add_docx_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    _style_table(table, widths)


def create_docx(product: dict[str, Any], path: Path) -> None:
    doc = Document()
    _configure_docx(doc, product)
    _add_docx_title(doc, product)
    doc.add_heading(SECTIONS[0], level=1)
    doc.add_paragraph(f"{product['model']} 的关键识别信息如下。相似型号参数不得混用。")
    _add_docx_table(
        doc,
        ["项目", "本型号参数"],
        [
            ["型号与识别", product["identity"]],
            ["供电规格", product["power"]],
            ["温度范围", product["temperature"]],
            ["兼容能力", product["compat"]],
            ["日志与存储", product["spec_detail"]],
        ],
        [2200, 7160],
    )
    doc.add_page_break()
    doc.add_heading(SECTIONS[1], level=1)
    for index, step in enumerate(
        [
            product["setup"],
            "完成首次设置后检查系统更新，再同步个人数据",
            product["feature"],
            product["daily_use"],
            product["backup"],
        ], 1
    ):
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(step.rstrip("。") + "。")
    doc.add_heading(SECTIONS[2], level=1)
    doc.add_paragraph(product["power"] + "。")
    doc.add_paragraph(product["temperature"] + "。")
    doc.add_paragraph(product["compat"] + "。")
    doc.add_paragraph(product["spec_detail"] + "。")
    doc.add_page_break()
    doc.add_heading(SECTIONS[3], level=1)
    doc.add_paragraph("维护与数据操作应按下表执行，避免配置尚未写入时断电。")
    _add_docx_table(
        doc,
        ["维护项目", "执行要求"],
        [["例行检查", product["maintenance"]], ["备份与恢复", product["backup"]]],
        [2600, 6760],
    )
    doc.add_heading(SECTIONS[4], level=1)
    _add_docx_table(
        doc,
        ["现象", "判断与处理"],
        [
            ["状态指示异常", product["indicator"]],
            ["配置写入提示", product["indicator_secondary"]],
            ["设备无响应", product["reset"]],
            ["连续启动失败", product["diagnostic"]],
        ],
        [2600, 6760],
    )
    doc.add_page_break()
    doc.add_heading(SECTIONS[5], level=1)
    doc.add_paragraph(product["warranty"] + "。")
    doc.add_paragraph(product["service"] + "。")
    doc.add_paragraph("保修判定以检测记录为准，寄送前应移除个人账户、存储卡和可拆卸附件。")
    doc.add_heading(SECTIONS[6], level=1)
    warning = doc.add_paragraph()
    warning.paragraph_format.space_before = Pt(8)
    warning.paragraph_format.space_after = Pt(8)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFF4E5")
    warning._p.get_or_add_pPr().append(shd)
    _font(warning.add_run("安全警告："), 11, True, "9B1C1C")
    _font(warning.add_run(product["safety"] + "。"), 11, False, "000000")
    doc.add_paragraph("发现异味、异常发热、外壳变形或液体侵入时，应立即断电并停止继续测试。")
    doc.add_paragraph("如仍无法解决，请记录状态灯、错误码和操作步骤后联系人工支持。")
    doc.save(path)


def _register_pdf_fonts() -> tuple[str, str]:
    regular = Path("C:/Windows/Fonts/msyh.ttc")
    bold = Path("C:/Windows/Fonts/msyhbd.ttc")
    if regular.exists():
        pdfmetrics.registerFont(TTFont("EvalCN", str(regular)))
        pdfmetrics.registerFont(TTFont("EvalCN-Bold", str(bold if bold.exists() else regular)))
        return "EvalCN", "EvalCN-Bold"
    return "Helvetica", "Helvetica-Bold"


def _create_pdf_platypus(product: dict[str, Any], path: Path) -> None:
    regular, bold = _register_pdf_fonts()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("EvalBody", parent=styles["BodyText"], fontName=regular, fontSize=10.5, leading=16, spaceAfter=8)
    h1 = ParagraphStyle("EvalH1", parent=styles["Heading1"], fontName=bold, fontSize=16, textColor=colors.HexColor("#2E74B5"), leading=21, spaceBefore=10, spaceAfter=10)
    title = ParagraphStyle("EvalTitle", parent=styles["Title"], fontName=bold, fontSize=27, leading=33, textColor=colors.HexColor("#0B2545"), alignment=TA_LEFT, spaceAfter=6)
    meta = ParagraphStyle("EvalMeta", parent=body, fontSize=9, textColor=colors.HexColor("#667085"), spaceAfter=18)
    warning = ParagraphStyle("EvalWarning", parent=body, textColor=colors.HexColor("#9B1C1C"), borderColor=colors.HexColor("#E6B8AF"), borderWidth=0.7, borderPadding=8, backColor=colors.HexColor("#FFF4E5"))

    base_path = path.with_suffix(".base.pdf")
    doc = BaseDocTemplate(str(base_path), pagesize=LETTER, leftMargin=inch, rightMargin=inch, topMargin=0.85 * inch, bottomMargin=0.8 * inch)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body")
    doc.addPageTemplates([PageTemplate(id="manual", frames=[frame])])
    table_style = TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), regular), ("FONTNAME", (0, 0), (-1, 0), bold),
        ("FONTSIZE", (0, 0), (-1, -1), 9), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#98A2B3")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ])
    story: list[Any] = [
        Paragraph("产品使用与服务指南", meta), Paragraph(product["model"], title),
        Paragraph(f"版本 {VERSION} | 合成评测文档，不代表真实商品规格", meta),
        Paragraph(SECTIONS[0], h1), Paragraph(f"{product['model']} 的关键识别信息如下。相似型号参数不得混用。", body),
        Table(
            [[Paragraph("项目", body), Paragraph("本型号参数", body)],
             [Paragraph("型号与识别", body), Paragraph(product["identity"], body)],
             [Paragraph("供电规格", body), Paragraph(product["power"], body)],
             [Paragraph("温度范围", body), Paragraph(product["temperature"], body)]],
            colWidths=[1.5 * inch, 5 * inch], style=table_style,
        ),
        PageBreak(), Paragraph(SECTIONS[1], h1),
        Paragraph("1. " + product["setup"] + "。", body),
        Paragraph("2. 完成首次设置后检查系统更新，再同步个人数据。", body),
        Paragraph("3. " + product["feature"] + "。", body),
        Paragraph(SECTIONS[2], h1), Paragraph(product["power"] + "。", body),
        Paragraph(product["temperature"] + "。", body), Paragraph(product["compat"] + "。", body),
        PageBreak(), Paragraph(SECTIONS[3], h1),
        Table(
            [[Paragraph("现象", body), Paragraph("判断与处理", body)],
             [Paragraph("状态指示异常", body), Paragraph(product["indicator"], body)],
             [Paragraph("设备无响应", body), Paragraph(product["reset"], body)]],
            colWidths=[1.8 * inch, 4.7 * inch], style=table_style,
        ),
        Paragraph(SECTIONS[4], h1), Paragraph(product["warranty"] + "。", body),
        Paragraph("安全警告：" + product["safety"] + "。", warning),
        Paragraph("如仍无法解决，请记录状态灯、错误码和操作步骤后联系人工支持。", body),
    ]
    doc.build(story)
    # 逐页叠加固定页眉页脚，避开部分 ReportLab/Poppler 组合在偶数页丢失
    # canvas 回调文本的问题，同时形成可检测的重复噪声。
    reader = PdfReader(str(base_path))
    writer = PdfWriter()
    for page_number, page in enumerate(reader.pages, 1):
        packet = BytesIO()
        overlay_canvas = Canvas(packet, pagesize=LETTER)
        overlay_canvas.setFont("Helvetica", 8.5)
        overlay_canvas.setFillColor(colors.HexColor("#667085"))
        overlay_canvas.drawString(inch, 10.45 * inch, f"EVAL HEADER | {product['model']} USER GUIDE")
        overlay_canvas.drawRightString(7.5 * inch, 0.55 * inch, f"EVAL FOOTER | {VERSION} | PAGE {page_number}")
        overlay_canvas.save()
        packet.seek(0)
        page.merge_page(PdfReader(packet).pages[0])
        writer.add_page(page)
    with path.open("wb") as stream:
        writer.write(stream)
    base_path.unlink()


def _canvas_lines(text: str, width: float, font: str, size: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if current and pdfmetrics.stringWidth(candidate, font, size) > width:
            lines.append(current)
            current = char
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def _canvas_text(
    canvas: Canvas, text: str, x: float, y: float, width: float,
    *, font: str, size: float = 10.5, leading: float = 16, color: str = "#000000",
) -> float:
    canvas.setFont(font, size)
    canvas.setFillColor(colors.HexColor(color))
    for line in _canvas_lines(text, width, font, size):
        canvas.drawString(x, y, line)
        y -= leading
    return y


def _canvas_header_footer(
    canvas: Canvas, product: dict[str, Any], page_number: int, font: str
) -> None:
    canvas.setFont(font, 8.5)
    canvas.setFillColor(colors.HexColor("#667085"))
    header = f"EVAL HEADER | {product['model']} USER GUIDE"
    footer = f"EVAL FOOTER | {VERSION} | PAGE {page_number}"

    def draw_chars(text: str, x: float, y: float) -> None:
        for char in text:
            canvas.drawString(x, y, char)
            x += pdfmetrics.stringWidth(char, font, 8.5)

    draw_chars(header, inch, 10.45 * inch)
    footer_width = sum(pdfmetrics.stringWidth(char, font, 8.5) for char in footer)
    draw_chars(footer, 7.5 * inch - footer_width, 0.55 * inch)


def _canvas_heading(canvas: Canvas, text: str, y: float, bold_font: str) -> float:
    canvas.setFont(bold_font, 16)
    canvas.setFillColor(colors.HexColor("#2E74B5"))
    canvas.drawString(inch, y, text)
    return y - 28


def _canvas_table(
    canvas: Canvas, x: float, y: float, widths: list[float], headers: list[str],
    rows: list[list[str]], regular: str, bold: str,
) -> float:
    all_rows = [headers, *rows]
    heights: list[float] = []
    for row in all_rows:
        line_count = max(
            len(_canvas_lines(value, width - 14, regular, 9.2))
            for value, width in zip(row, widths, strict=True)
        )
        heights.append(max(28, line_count * 13 + 12))
    total_width = sum(widths)
    cursor_y = y
    canvas.setStrokeColor(colors.HexColor("#98A2B3"))
    canvas.setLineWidth(0.5)
    for row_index, (row, height) in enumerate(zip(all_rows, heights, strict=True)):
        if row_index == 0:
            canvas.setFillColor(colors.HexColor("#E8EEF5"))
            canvas.rect(x, cursor_y - height, total_width, height, stroke=0, fill=1)
        cursor_x = x
        for value, width in zip(row, widths, strict=True):
            canvas.rect(cursor_x, cursor_y - height, width, height, stroke=1, fill=0)
            lines = _canvas_lines(value, width - 14, regular, 9.2)
            text_y = cursor_y - 11 - (height - len(lines) * 13) / 2
            canvas.setFont(bold if row_index == 0 else regular, 9.2)
            canvas.setFillColor(colors.black)
            for line in lines:
                canvas.drawString(cursor_x + 7, text_y, line)
                text_y -= 13
            cursor_x += width
        cursor_y -= height
    return cursor_y


def create_pdf(product: dict[str, Any], path: Path) -> None:
    """使用固定坐标生成四页 PDF，保证表格与分页在 Poppler 下稳定。"""

    regular, bold = _register_pdf_fonts()
    canvas = Canvas(str(path), pagesize=LETTER)
    content_width = 6.5 * inch

    y = 9.72 * inch
    y = _canvas_text(canvas, "产品使用与服务指南", inch, y, content_width, font=regular, size=10, leading=15, color="#667085")
    y -= 14
    y = _canvas_text(canvas, product["model"], inch, y, content_width, font=bold, size=27, leading=32, color="#0B2545")
    y = _canvas_text(canvas, f"版本 {VERSION} | 合成评测文档，不代表真实商品规格", inch, y - 4, content_width, font=regular, size=9.5, leading=16, color="#667085")
    y = _canvas_heading(canvas, SECTIONS[0], y - 20, bold)
    y = _canvas_text(canvas, f"{product['model']} 的关键识别信息如下。相似型号参数不得混用。", inch, y, content_width, font=regular)
    _canvas_table(
        canvas, inch, y - 8, [1.5 * inch, 5 * inch], ["项目", "本型号参数"],
        [
            ["型号与识别", product["identity"]],
            ["供电规格", product["power"]],
            ["温度范围", product["temperature"]],
            ["兼容能力", product["compat"]],
            ["日志与存储", product["spec_detail"]],
        ],
        regular, bold,
    )
    canvas.showPage()

    y = _canvas_heading(canvas, SECTIONS[1], 9.72 * inch, bold)
    for index, value in enumerate(
        [
            product["setup"],
            "完成首次设置后检查系统更新，再同步个人数据",
            product["feature"],
            product["daily_use"],
            product["backup"],
        ], 1
    ):
        y = _canvas_text(canvas, f"{index}. {value}。", inch, y, content_width, font=regular)
        y -= 6
    y = _canvas_heading(canvas, SECTIONS[2], y - 5, bold)
    for value in (product["power"], product["temperature"], product["compat"], product["spec_detail"]):
        y = _canvas_text(canvas, value + "。", inch, y, content_width, font=regular)
        y -= 7
    canvas.showPage()

    y = _canvas_heading(canvas, SECTIONS[3], 9.72 * inch, bold)
    y = _canvas_text(canvas, "维护与数据操作应按下表执行，避免配置尚未写入时断电。", inch, y, content_width, font=regular)
    y = _canvas_table(
        canvas, inch, y - 6, [1.8 * inch, 4.7 * inch], ["维护项目", "执行要求"],
        [["例行检查", product["maintenance"]], ["备份与恢复", product["backup"]]], regular, bold,
    )
    y = _canvas_heading(canvas, SECTIONS[4], y - 20, bold)
    y = _canvas_table(
        canvas, inch, y, [1.8 * inch, 4.7 * inch], ["现象", "判断与处理"],
        [
            ["状态指示异常", product["indicator"]],
            ["配置写入提示", product["indicator_secondary"]],
            ["设备无响应", product["reset"]],
            ["连续启动失败", product["diagnostic"]],
        ], regular, bold,
    )
    canvas.showPage()

    y = _canvas_heading(canvas, SECTIONS[5], 9.72 * inch, bold)
    y = _canvas_text(canvas, product["warranty"] + "。", inch, y, content_width, font=regular)
    y = _canvas_text(canvas, product["service"] + "。", inch, y - 8, content_width, font=regular)
    y = _canvas_text(canvas, "保修判定以检测记录为准，寄送前应移除个人账户、存储卡和可拆卸附件。", inch, y - 8, content_width, font=regular)
    y = _canvas_heading(canvas, SECTIONS[6], y - 20, bold)
    warning_lines = _canvas_lines("安全警告：" + product["safety"] + "。", content_width - 16, regular, 10.5)
    warning_height = len(warning_lines) * 16 + 16
    canvas.setFillColor(colors.HexColor("#FFF4E5"))
    canvas.setStrokeColor(colors.HexColor("#E6B8AF"))
    canvas.rect(inch - 4, y - warning_height + 4, content_width + 8, warning_height, stroke=1, fill=1)
    y = _canvas_text(canvas, "安全警告：" + product["safety"] + "。", inch + 4, y - 7, content_width - 8, font=regular, color="#9B1C1C")
    y = _canvas_text(canvas, "发现异味、异常发热、外壳变形或液体侵入时，应立即断电并停止继续测试。", inch, y - 18, content_width, font=regular)
    _canvas_text(canvas, "如仍无法解决，请记录状态灯、错误码和操作步骤后联系人工支持。", inch, y - 8, content_width, font=regular)
    canvas.save()


def create_markdown(product: dict[str, Any], path: Path) -> None:
    content = f"""# {product['model']} 产品使用与服务指南

> 版本 {VERSION}。合成评测文档，不代表真实商品规格。

## {SECTIONS[0]}

{product['model']} 的关键识别信息如下。相似型号参数不得混用。

| 项目 | 本型号参数 |
| --- | --- |
| 型号与识别 | {product['identity']} |
| 供电规格 | {product['power']} |
| 温度范围 | {product['temperature']} |
| 兼容能力 | {product['compat']} |
| 日志与存储 | {product['spec_detail']} |

## {SECTIONS[1]}

1. {product['setup']}。
2. 完成首次设置后检查系统更新，再同步个人数据。
3. {product['feature']}。
4. {product['daily_use']}。
5. {product['backup']}。

## {SECTIONS[2]}

{product['power']}。

{product['temperature']}。

{product['compat']}。

{product['spec_detail']}。

## {SECTIONS[3]}

维护与数据操作应按下表执行，避免配置尚未写入时断电。

| 维护项目 | 执行要求 |
| --- | --- |
| 例行检查 | {product['maintenance']} |
| 备份与恢复 | {product['backup']} |

## {SECTIONS[4]}

| 现象 | 判断与处理 |
| --- | --- |
| 状态指示异常 | {product['indicator']} |
| 配置写入提示 | {product['indicator_secondary']} |
| 设备无响应 | {product['reset']} |
| 连续启动失败 | {product['diagnostic']} |

## {SECTIONS[5]}

{product['warranty']}。

{product['service']}。

保修判定以检测记录为准，寄送前应移除个人账户、存储卡和可拆卸附件。

## {SECTIONS[6]}

**安全警告：{product['safety']}。**

发现异味、异常发热、外壳变形或液体侵入时，应立即断电并停止继续测试。

如仍无法解决，请记录状态灯、错误码和操作步骤后联系人工支持。
"""
    path.write_text(content, encoding="utf-8")


def _page_for(product: dict[str, Any], section: str) -> int | None:
    if product["format"] != "pdf":
        return None
    if section == SECTIONS[0]:
        return 1
    if section in SECTIONS[1:3]:
        return 2
    if section in SECTIONS[3:5]:
        return 3
    return 4


def make_questions() -> list[dict[str, Any]]:
    questions: list[dict[str, Any]] = []
    templates = [
        ("identity", "{model} 的屏幕/形态和识别码或硬件型号是什么？", SECTIONS[0], "identity", "table_extraction"),
        ("power", "{model} 的供电或充电功率与端口规则是什么？", SECTIONS[0], "power", "table_extraction"),
        ("temperature", "{model} 允许或建议的充电、工作温度范围是多少？", SECTIONS[0], "temperature", "exact_numeric"),
        ("setup", "{model} 首次设置时第一步应该怎么做？", SECTIONS[1], "setup", "ordered_list"),
        ("feature", "{model} 的特殊模式有什么启用条件或限制？", SECTIONS[1], "feature", "section_boundary"),
        ("indicator", "{model} 出现文档中的特定状态灯或错误码代表什么？", SECTIONS[4], "indicator", "troubleshooting_table"),
        ("reset", "{model} 无响应时应如何按文档执行复位或强制重启？", SECTIONS[4], "reset", "troubleshooting_table"),
        ("compat", "{model} 的协议、端口或网络兼容性限制是什么？", SECTIONS[2], "compat", "hard_negative"),
        ("warranty", "{model} 的保修期限或主要除外条件是什么？", SECTIONS[5], "warranty", "policy_boundary"),
        ("safety", "使用 {model} 时文档明确禁止或要求的安全事项是什么？", SECTIONS[6], "safety", "warning_block"),
    ]
    number = 1
    for product in PRODUCTS:
        for suffix, query, section, field, case_type in templates:
            fact = product[field]
            standard_answer = fact if fact.endswith(("。", "！", "？")) else f"{fact}。"
            questions.append({
                "id": f"retrieval_{number:03d}",
                "query": query.format(model=product["model"]),
                "standard_answer": standard_answer,
                "case_type": case_type,
                "filters": {"document_type": "product_manual", "product_model": product["id"], "category": product["category"]},
                "relevant_targets": [{
                    "document_id": product["id"], "section": section,
                    "page_start": _page_for(product, section),
                    "expected_rendered_page": (
                        1 if section == SECTIONS[0]
                        else 2 if section in SECTIONS[1:3]
                        else 3 if section in SECTIONS[3:5]
                        else 4
                    ),
                    "evidence_fragments": [fact],
                }],
                "required_facts": [fact],
                "hard_negative_document_ids": [product["sibling"]],
            })
            number += 1
    return questions


def write_qa_reference() -> None:
    """生成适合人工抽查的问题—标准答案清单。"""

    questions = make_questions()
    models = {product["id"]: product["model"] for product in PRODUCTS}
    lines = [
        f"# RAG 检索评测：{len(questions)} 道问题与标准答案",
        "",
        "> 所有产品与参数均为虚构测试数据。标准答案必须能由标注的证据切片直接支持。",
        "",
    ]
    current_document = None
    ordinal = 0
    for question in questions:
        target = question["relevant_targets"][0]
        document_id = target["document_id"]
        if document_id != current_document:
            current_document = document_id
            ordinal = 0
            lines.extend([f"## {models[document_id]}（`{document_id}`）", ""])
        ordinal += 1
        lines.extend([
            f"### {ordinal}. {question['query']}",
            "",
            f"**标准答案：** {question['standard_answer']}",
            "",
            f"**标准位置：** `{document_id}::{target['section']}`",
            "",
            f"**难负例：** `{question['hard_negative_document_ids'][0]}`",
            "",
        ])
    (OUTPUT_ROOT / "qa_reference.md").write_text("\n".join(lines), encoding="utf-8")


def build_metadata() -> None:
    manifest_documents = []
    expectations = []
    for product in PRODUCTS:
        filename = f"{product['id']}_guide.{product['format']}"
        manifest_documents.append({
            "document_id": product["id"], "path": f"documents/{filename}",
            "document_type": "product_manual", "product_model": product["id"],
            "display_model": product["model"], "category": product["category"],
            "version": VERSION, "format": product["format"],
        })
        expectations.append({
            "document_id": product["id"],
            "expected_headings": SECTIONS,
            "expected_tables": [
                {"section": SECTIONS[0], "headers": ["项目", "本型号参数"], "data_row_count": 5},
                {"section": SECTIONS[3], "headers": ["维护项目", "执行要求"], "data_row_count": 2},
                {"section": SECTIONS[4], "headers": ["现象", "判断与处理"], "data_row_count": 4},
            ],
            "expected_numbered_list_items": 5,
            "required_content_fragments": [
                product[key]
                for key in (
                    "identity", "power", "temperature", "setup", "feature", "indicator",
                    "reset", "compat", "warranty", "safety", "spec_detail", "daily_use",
                    "maintenance", "backup", "indicator_secondary", "diagnostic", "service",
                )
            ],
            "expected_removed_noise": (
                [f"星澜数字设备支持中心 | {product['model']} 用户指南", f"内部评测样本 | {VERSION}"]
                if product["format"] == "docx" else []
            ),
            "expected_rendered_page_count": 4 if product["format"] in {"docx", "pdf"} else None,
        })
    (OUTPUT_ROOT / "manifest.json").write_text(
        json.dumps({"benchmark_version": VERSION, "documents": manifest_documents}, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUTPUT_ROOT / "questions.json").write_text(
        json.dumps({"benchmark_version": VERSION, "questions": make_questions()}, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUTPUT_ROOT / "cleaning_expectations.json").write_text(
        json.dumps({"benchmark_version": VERSION, "documents": expectations}, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def write_readme() -> None:
    text = """# 证据级 RAG 清洗与切片评测集 v2

本目录包含 50 份完全虚构的设备文档和 500 道证据级查询。所有参数只用于测试，
不得用于真实商品答复。型号成对设计为 hard negatives，避免只靠品类词完成召回。

## 文件

- `manifest.json`：入库元数据与源文件路径。
- `questions.json`：500 道问题、标准答案、正确文档、章节、PDF 页码、证据和相似型号负例。
- `qa_reference.md`：便于人工抽查的 500 道问题与标准答案清单。
- `cleaning_expectations.json`：清洗阶段应保留的标题、表格、列表、正文及应删除的页眉页脚。
- `documents/`：混合 DOCX、PDF 和 Markdown，共 50 份。

## 推荐指标

清洗评测：必需正文保留率、标题恢复率、表格恢复率、列表恢复率、页眉页脚残留率、
重复文本率和乱码率。切片评测：证据 Chunk Recall@1/3/5、MRR、nDCG@5、完整事实
同块率、跨章节污染率、相似型号泄漏率，以及 Top-K 证据覆盖率。

DOCX 的源格式没有稳定页码，所以 `page_start` 为 null，`expected_rendered_page` 只供
渲染核验。PDF 的 `page_start` 是可自动校验的真实页码。

## 运行检索评测

下面的命令复用项目正式的清洗、切块、Dense、BM25、RRF 和重排逻辑，不调用生成 LLM：

```powershell
.\\.venv\\Scripts\\python.exe .\\agent\\scripts\\evaluate_retrieval_eval_v1.py
```

默认同时执行“不使用型号过滤”和“使用完整元数据过滤”两轮，结果写入
`agent/runtime/retrieval_eval_v1_report.json` 与同名 Markdown 文件。可使用
`--max-questions 10` 快速试跑。
"""
    (OUTPUT_ROOT / "README.md").write_text(text, encoding="utf-8")


def main() -> None:
    resolved_output = OUTPUT_ROOT.resolve()
    if ROOT.resolve() not in resolved_output.parents:
        raise RuntimeError(f"拒绝清理工作区之外的路径：{resolved_output}")
    if OUTPUT_ROOT.exists():
        shutil.rmtree(OUTPUT_ROOT)
    DOCUMENT_ROOT.mkdir(parents=True, exist_ok=True)
    for product in PRODUCTS:
        path = DOCUMENT_ROOT / f"{product['id']}_guide.{product['format']}"
        if product["format"] == "docx":
            create_docx(product, path)
        elif product["format"] == "pdf":
            create_pdf(product, path)
        else:
            create_markdown(product, path)
    build_metadata()
    write_qa_reference()
    write_readme()
    print(f"generated={OUTPUT_ROOT} documents={len(PRODUCTS)} questions={len(make_questions())}")


if __name__ == "__main__":
    main()
