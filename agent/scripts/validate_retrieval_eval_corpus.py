"""验证 retrieval_eval_v1 的源文件、清洗期望和证据级切片标注。"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document as WordDocument
from docx.oxml.ns import qn


AGENT_ROOT = Path(__file__).resolve().parents[1]
if str(AGENT_ROOT) not in sys.path:
    sys.path.insert(0, str(AGENT_ROOT))

from core.knowledge.ingestion import DocumentPreprocessor, StructuredDocumentChunker  # noqa: E402


ROOT = AGENT_ROOT / "data" / "knowledge" / "retrieval_eval_v1"
VERSION_TEXT = "2026.08-v2"


def _norm(text: str) -> str:
    return re.sub(r"\s+|[。；，、,.|]", "", text).casefold()


def main() -> None:
    manifest = json.loads((ROOT / "manifest.json").read_text(encoding="utf-8"))
    questions = json.loads((ROOT / "questions.json").read_text(encoding="utf-8"))["questions"]
    expectations = {
        item["document_id"]: item
        for item in json.loads((ROOT / "cleaning_expectations.json").read_text(encoding="utf-8"))["documents"]
    }
    errors: list[str] = []
    chunks_by_document: dict[str, list[object]] = {}
    parsed_summary: list[dict[str, object]] = []
    parser = DocumentPreprocessor()
    chunker = StructuredDocumentChunker()

    for item in manifest["documents"]:
        source = ROOT / item["path"]
        if not source.is_file():
            errors.append(f"missing source: {source}")
            continue
        parsed = parser.parse(source, document_id=item["document_id"])
        chunks = chunker.chunk(
            parsed,
            document_type=item["document_type"],
            product_model=item["product_model"],
            category=item["category"],
            version=item["version"],
        )
        chunks_by_document[item["document_id"]] = chunks
        all_text = _norm("\n".join(block.text for block in parsed.blocks))
        paths = {block.section_path for block in parsed.blocks if block.section_path}
        expectation = expectations[item["document_id"]]
        for fragment in expectation["required_content_fragments"]:
            if _norm(fragment) not in all_text:
                errors.append(f"{item['document_id']}: missing cleaned fragment: {fragment}")
        for heading in expectation["expected_headings"]:
            if not any(heading in path for path in paths):
                errors.append(f"{item['document_id']}: missing heading: {heading}; got={sorted(paths)}")
        for noise in expectation["expected_removed_noise"]:
            if _norm(noise) in all_text:
                errors.append(f"{item['document_id']}: header/footer noise remains: {noise}")
        parsed_tables = [block.table for block in parsed.blocks if block.table is not None]
        for expected_table in expectation["expected_tables"]:
            matching_tables = [
                table for table in parsed_tables
                if all(header in table.headers for header in expected_table["headers"])
            ]
            if not matching_tables:
                errors.append(f"{item['document_id']}: expected table missing: {expected_table['headers']}")
            elif len(matching_tables[0].rows) < expected_table["data_row_count"]:
                errors.append(
                    f"{item['document_id']}: table rows {len(matching_tables[0].rows)} "
                    f"< {expected_table['data_row_count']}"
                )

        if item["format"] == "docx":
            word = WordDocument(source)
            section = word.sections[0]
            geometry = [
                section.page_width.inches, section.page_height.inches,
                section.top_margin.inches, section.right_margin.inches,
                section.bottom_margin.inches, section.left_margin.inches,
            ]
            expected_geometry = [8.5, 11.0, 1.0, 1.0, 1.0, 1.0]
            if any(abs(actual - expected) > 0.01 for actual, expected in zip(geometry, expected_geometry, strict=True)):
                errors.append(f"{item['document_id']}: invalid DOCX page geometry: {geometry}")
            if len(word._element.xpath('.//w:br[@w:type="page"]')) != 3:
                errors.append(f"{item['document_id']}: expected exactly three page breaks")
            if len(word.tables) != 3:
                errors.append(f"{item['document_id']}: expected three DOCX tables")
            if sum(paragraph.style.name == "List Number" for paragraph in word.paragraphs) != 5:
                errors.append(f"{item['document_id']}: expected five real numbered-list paragraphs")
            header_text = "\n".join(paragraph.text for paragraph in section.header.paragraphs)
            footer_text = "\n".join(paragraph.text for paragraph in section.footer.paragraphs)
            if item["display_model"] not in header_text or VERSION_TEXT not in footer_text:
                errors.append(f"{item['document_id']}: DOCX header/footer content missing")
            for table in word.tables:
                tbl_pr = table._tbl.tblPr
                tbl_w = tbl_pr.first_child_found_in("w:tblW")
                ind = tbl_pr.first_child_found_in("w:tblInd")
                widths = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
                if tbl_w is None or int(tbl_w.get(qn("w:w"))) != 9360 or sum(widths) != 9360:
                    errors.append(f"{item['document_id']}: DOCX table width/grid mismatch")
                if ind is None or int(ind.get(qn("w:w"))) != 120:
                    errors.append(f"{item['document_id']}: DOCX table indent is not 120 DXA")
        parsed_summary.append(
            {
                "document_id": item["document_id"],
                "format": item["format"],
                "block_count": len(parsed.blocks),
                "chunk_count": len(chunks),
                "table_count": sum(block.table is not None for block in parsed.blocks),
                "heading_paths": sorted(paths),
            }
        )

    if len(questions) != 500:
        errors.append(f"expected 500 questions, got {len(questions)}")
    if len({item["id"] for item in questions}) != len(questions):
        errors.append("duplicate question ids")
    for question in questions:
        standard_answer = question.get("standard_answer", "")
        if not standard_answer:
            errors.append(f"{question['id']}: standard_answer missing")
        elif not all(_norm(fact) in _norm(standard_answer) for fact in question.get("required_facts", [])):
            errors.append(f"{question['id']}: standard_answer does not cover required facts")
        for target in question["relevant_targets"]:
            chunks = chunks_by_document.get(target["document_id"], [])
            matching = [
                chunk for chunk in chunks
                if target["section"] in (chunk.section_path or chunk.section)
            ]
            if not matching:
                errors.append(
                    f"{question['id']}: no chunk for {target['document_id']}::{target['section']}"
                )
                continue
            evidence = _norm("\n".join(chunk.content for chunk in matching))
            for fragment in target["evidence_fragments"]:
                if _norm(fragment) not in evidence:
                    errors.append(f"{question['id']}: evidence missing from target chunks: {fragment}")
            page = target.get("page_start")
            if page is not None and not any(
                chunk.page_start is not None
                and chunk.page_start <= page <= (chunk.page_end or chunk.page_start)
                for chunk in matching
            ):
                errors.append(f"{question['id']}: expected PDF page {page} not represented")

    report = {
        "document_count": len(manifest["documents"]),
        "question_count": len(questions),
        "chunk_count": sum(len(chunks) for chunks in chunks_by_document.values()),
        "error_count": len(errors),
        "documents": parsed_summary,
        "errors": errors,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
